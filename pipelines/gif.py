"""デジタル庁 政府相互運用性フレームワーク（GIF）の語彙を取り出す。

GIF は GitHub（digital-go-jp/GIF）でタグを切って版を出している。最新のタグを 1 つ決め、
その版の Excel 3 つを CSV にする。版は各行の version 列に入れる。

出力:
- data/gif/term.csv（1行 = コアデータモデルの 1 項目。438_コアデータモデル_DMD.xlsx）
- data/gif/code.csv（1行 = コード一覧の 1 値。438-1_コアデータモデル_コード一覧.xlsx）
- data/gif/code_system.csv（1行 = 行政が使うコード体系 1 つ。491-2_コード_コード一覧_*.xlsx）

データソース: デジタル庁 https://github.com/digital-go-jp/GIF
"""

import csv
import logging
import re
import subprocess
import tempfile
from functools import reduce
from pathlib import Path

import openpyxl

logger = logging.getLogger("pipelines")

REPO = "digital-go-jp/GIF"
OUTPUT_DIR = Path("data/gif")

# 取り込む Excel。491-2 はファイル名に日付が入って変わるので、接頭辞で探す
FILES = {
    "term": "430_コアデータモデル/438_",
    "code": "430_コアデータモデル/438-1_",
    "code_system": "490_その他/491-2_",
}


def latest_tag() -> str:
    out = subprocess.run(
        ["git", "ls-remote", "--tags", "--refs", f"https://github.com/{REPO}.git"],
        check=True, capture_output=True, text=True,
    ).stdout
    tags = [line.split("refs/tags/")[1] for line in out.splitlines()]
    return max(tags, key=lambda t: tuple(int(n) for n in re.findall(r"\d+", t)))


def checkout(tag: str, dest: Path) -> Path:
    # Excel だけ欲しいが、ファイル名に日付が入るものがあるので木ごと取って探す
    subprocess.run(
        ["git", "-c", "advice.detachedHead=false", "clone", "--quiet", "--depth", "1", "--branch", tag, f"https://github.com/{REPO}.git", str(dest)],
        check=True,
    )
    return dest


def find_paths(root: Path) -> dict[str, Path]:
    def pick(prefix: str) -> Path:
        found = sorted(root.glob(f"{prefix}*.xlsx"))
        if len(found) != 1:
            raise SystemExit(f"{REPO}: {prefix}*.xlsx が {len(found)} 個ある")
        return found[0]

    return {name: pick(prefix) for name, prefix in FILES.items()}


def clean(value) -> str:
    # 項目名には「事業所住所 (*1)」のような注記の印が付く
    return re.sub(r"\s*\(\*\d+\)$", "", "" if value is None else str(value).strip())


def header_index(row, headers: dict[str, str]) -> dict[str, int]:
    names = [clean(c).replace("\n", "") for c in row]
    return {headers[n]: i for i, n in enumerate(names) if n in headers}


# --- 438_コアデータモデル_DMD.xlsx -----------------------------------------
# シートがクラス（個人・法人・住所…）。「基礎項目」「拡張項目」の見出し行のあとに
# 「No.」で始まる表が続く。拡張項目の行には No. が無い

TERM_HEADERS = {
    "No.": "no",
    "項目名": "term_ja",
    "項目名（英語）": "term_en",
    "説明": "description",
    "形式": "format",
    "記入例": "example",
    "名称・形式等の参考元": "reference",
}


def term_rows(sheet) -> list[dict]:
    def step(state, row):
        columns, section, out = state
        first = row[0] if row else None
        if first == "No.":
            return header_index(row, TERM_HEADERS), section, out
        if isinstance(first, str) and first.startswith(("基礎項目", "拡張項目")):
            return columns, "base" if first.startswith("基礎") else "extension", out
        is_term = isinstance(first, int) or (first is None and section == "extension")
        if is_term and columns and row[columns["term_ja"]]:
            values = {k: clean(row[i]) for k, i in columns.items()}
            return columns, section, [*out, {"class_name": sheet.title, "section": section, **values}]
        return state

    return reduce(step, (tuple(r) for r in sheet.iter_rows(values_only=True)), ({}, "base", []))[2]


# --- 438-1_コアデータモデル_コード一覧.xlsx --------------------------------
# 1 シート。コード一覧の見出し行（No. と名称がある）の下に、値の行（Code と Value だけ）が続く

CODE_HEADERS = {
    "コード一覧No.": "code_list_no",
    "コード一覧名称": "code_list_name",
    "コード一覧名称（英語）": "code_list_name_en",
    "Code値（=コード）": "code",
    "Value値（=名称）": "value",
    "説明": "description",
}


def code_rows(sheet) -> list[dict]:
    def step(state, row):
        columns, current, out = state
        if not columns:
            found = header_index(row, CODE_HEADERS)
            return (found if "code" in found else {}), current, out
        cell = {k: clean(row[i]) for k, i in columns.items()}
        if cell["code_list_no"]:
            return columns, cell, out
        if current and cell["code"]:
            heading = {k: current[k] for k in ("code_list_no", "code_list_name", "code_list_name_en")}
            return columns, current, [*out, {**heading, "code": cell["code"], "value": cell["value"], "description": cell["description"]}]
        return state

    return reduce(step, (tuple(r) for r in sheet.iter_rows(values_only=True)), ({}, None, []))[2]


# --- 491-2_コード_コード一覧_*.xlsx ---------------------------------------
# 「はじめに」以外のシートが分野。見出し行（通しNo で始まる）の位置はシートごとに違う。
# 通しNo は =ROW()-1 という式で値を持たないので、シートの中の並び順で振り直す

CODE_SYSTEM_HEADERS = {
    "通しNo": "no",
    "分類": "category",
    "名前": "name",
    "データのURL": "data_url",
    "データ形式": "data_format",
    "説明": "description",
    "公開ウェブページ": "web_page",
    "公表組織名": "publisher",
}


def code_system_rows(sheet) -> list[dict]:
    def step(state, row):
        columns, out = state
        if not columns:
            found = header_index(row, CODE_SYSTEM_HEADERS)
            return (found if "name" in found else {}), out
        if row[columns["no"]] is not None and row[columns["name"]]:
            values = {k: clean(row[i]) for k, i in columns.items() if k != "no"}
            return columns, [*out, {"field": sheet.title, "no": len(out) + 1, **values}]
        return state

    return reduce(step, (tuple(r) for r in sheet.iter_rows(values_only=True)), ({}, []))[1]


EXTRACT = {
    "term": lambda book: [r for s in book.worksheets for r in term_rows(s)],
    "code": lambda book: [r for s in book.worksheets for r in code_rows(s)],
    "code_system": lambda book: [r for s in book.worksheets if s.title != "はじめに" for r in code_system_rows(s)],
}


def write_csv(path: Path, rows: list[dict]) -> None:
    fields = list(dict.fromkeys(k for r in rows for k in r))
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def download_gif(tag: str | None = None, output_dir: Path = OUTPUT_DIR) -> dict[str, int]:
    tag = tag or latest_tag()
    logger.info("GIF %s", tag)
    output_dir.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory() as tmp:
        root = checkout(tag, Path(tmp) / "GIF")
        paths = find_paths(root)

        def one(name: str) -> int:
            source = paths[name].relative_to(root).as_posix()
            book = openpyxl.load_workbook(paths[name], read_only=True)
            rows = [{"version": tag, "source_path": source, **r} for r in EXTRACT[name](book)]
            if not rows:
                raise SystemExit(f"{source}: 1 行も取れなかった（Excel の形が変わった？）")
            write_csv(output_dir / f"{name}.csv", rows)
            logger.info("  %s: %d 行", name, len(rows))
            return len(rows)

        return {name: one(name) for name in FILES}
